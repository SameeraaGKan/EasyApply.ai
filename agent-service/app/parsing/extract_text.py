import io

import docx
from pypdf import PdfReader


class UnsupportedFileTypeError(ValueError):
    pass


def extract_text(file_bytes: bytes, file_type: str) -> str:
    """Extracts raw text from a resume file. Supports the two formats the
    upload UI accepts (see web/src/app/(app)/resume/resume-upload-form.tsx).
    Raises UnsupportedFileTypeError for anything else."""
    if file_type == "pdf":
        return _extract_pdf_text(file_bytes)
    if file_type == "docx":
        return _extract_docx_text(file_bytes)
    raise UnsupportedFileTypeError(f"Unsupported file_type: {file_type}")


def _extract_pdf_text(file_bytes: bytes) -> str:
    reader = PdfReader(io.BytesIO(file_bytes))
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n".join(pages).strip()


def _extract_docx_text(file_bytes: bytes) -> str:
    document = docx.Document(io.BytesIO(file_bytes))
    paragraphs = [p.text for p in document.paragraphs]
    # Resumes frequently put contact info / skills in tables rather than
    # paragraphs — pull those in too, in reading order per table.
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    paragraphs.append(cell.text)
    return "\n".join(p for p in paragraphs if p.strip()).strip()
